from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path
from textwrap import wrap
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image as PDFImage,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DIAGRAMS = OUT / "diagramas"
FIXTURES = OUT / "fixtures"
TEMPLATE = Path.home() / "Downloads" / "Plant de tests de software para Calidad.xlsx"
DATE_LABEL = "27 de septiembre de 2026"
NAVY = "102B46"
INK = "19334C"
BLUE = "3465E9"
ORANGE = "F05A24"
MINT = "E8F3EF"
PAPER = "F4F3EF"
GRAY = "667085"
FONT_PATH = "/System/Library/Fonts/Supplemental/Arial.ttf"


PEOPLE = {
    "JhonRamirez22": "JhonRamirez22",
    "Miguel Felipe Ceballos Ramirez": "Miguel Felipe Ceballos Ramirez",
    "Johan Steven Muñoz Enriquez": "Johan Steven Muñoz Enriquez",
    "stheban hoyos villota": "stheban hoyos villota",
    "Supersamuel 909": "Supersamuel 909",
}

ISSUES = [
    {
        "key": "RTE-10", "hu": "HU-01", "title": "Registro de usuarios con roles",
        "actor": "administrador", "want": "registrar usuarios con roles (admin, gerente, operador, contador)",
        "purpose": "controlar quién tiene acceso a cada módulo del sistema", "assignee": "JhonRamirez22",
        "criteria": ["El formulario permite crear un usuario con nombre, email, contraseña y rol.", "Los roles disponibles son administrador, gerente, operador y contador.", "No se permiten emails duplicados.", "La contraseña cumple requisitos mínimos de seguridad: 8 o más caracteres, una mayúscula y un número.", "El usuario se guarda en estado activo por defecto.", "Se muestra confirmación exitosa tras el registro."],
    },
    {
        "key": "RTE-11", "hu": "HU-02", "title": "Autenticación con email y JWT",
        "actor": "usuario del sistema", "want": "autenticarme con email y contraseña utilizando JWT",
        "purpose": "acceder de forma segura y mantener mi sesión protegida", "assignee": "JhonRamirez22",
        "criteria": ["El login recibe email y contraseña.", "Las credenciales incorrectas retornan un error genérico que no revela si el email existe.", "El JWT expira en 24 horas y se renueva automáticamente.", "El token contiene el id, email y rol del usuario.", "El logout invalida el token en el cliente.", "La cuenta se bloquea tras cinco intentos fallidos."],
    },
    {
        "key": "RTE-15", "hu": "HU-06", "title": "Registro de productos con variedad y origen",
        "actor": "operador", "want": "registrar productos con variedad, origen y certificación",
        "purpose": "mantener un catálogo completo y confiable para la exportación", "assignee": "Miguel Felipe Ceballos Ramirez",
        "criteria": ["El formulario permite crear productos con nombre, tipo (café/cacao), variedad, origen geográfico, certificación y descripción.", "Se pueden asociar múltiples certificaciones a un producto (orgánico, Fair Trade, Rainforest, etc.).", "El producto queda activo por defecto.", "No se permiten productos duplicados con la misma variedad y origen.", "Se puede buscar y filtrar por tipo, variedad y estado."],
    },
    {
        "key": "RTE-16", "hu": "HU-07", "title": "Registro de lotes con trazabilidad",
        "actor": "operador", "want": "registrar lotes con número de trazabilidad único",
        "purpose": "cumplir las regulaciones europeas de trazabilidad de origen de café y cacao", "assignee": "Miguel Felipe Ceballos Ramirez",
        "criteria": ["Cada lote se asocia a un producto registrado.", "El número de trazabilidad se genera automáticamente o se ingresa manualmente.", "El lote incluye peso, fecha de cosecha, fecha de procesamiento y ubicación de origen.", "Se pueden registrar análisis de calidad asociados al lote.", "El lote tiene estados Disponible, Reservado, Enviado y Certificado.", "Se puede rastrear la trazabilidad completa desde origen hasta destino."],
    },
    {
        "key": "RTE-17", "hu": "HU-08", "title": "Gestión de certificados de calidad",
        "actor": "operador", "want": "subir y gestionar certificados de calidad (orgánico, Fair Trade, Rainforest Alliance)",
        "purpose": "respaldar la calidad y autenticidad de los productos exportados", "assignee": "Miguel Felipe Ceballos Ramirez",
        "criteria": ["Se pueden subir certificados PDF, JPG o PNG.", "Cada certificado se asocia a un producto o lote.", "El certificado incluye tipo, número, fecha de emisión y vencimiento, y entidad emisora.", "Se puede visualizar y descargar el adjunto.", "Los certificados próximos a vencer (30 días) se marcan con una alerta visual.", "Se puede eliminar un certificado obsoleto con confirmación."],
    },
    {
        "key": "RTE-18", "hu": "HU-09", "title": "Consulta de inventario por producto y lote",
        "actor": "gerente", "want": "visualizar el inventario completo desglosado por producto, lote y estado",
        "purpose": "tomar decisiones informadas sobre disponibilidad, pedidos y planificación de exportaciones", "assignee": "Miguel Felipe Ceballos Ramirez",
        "criteria": ["El dashboard muestra inventario total por producto y resumen de lotes.", "Se filtra por tipo (café/cacao), estado y rango de fechas.", "Cada lote muestra peso disponible, estado, ubicación y trazabilidad.", "El inventario se exporta a Excel o PDF.", "Los datos se actualizan en tiempo real al registrar movimientos de lote.", "Se muestra el porcentaje de disponibilidad por producto."],
    },
    {
        "key": "RTE-20", "hu": "HU-11", "title": "Registro de clientes europeos",
        "actor": "operador", "want": "registrar clientes europeos con empresa, dirección, país y VAT ID",
        "purpose": "mantener un directorio de clientes internacionales para la exportación", "assignee": "Johan Steven Muñoz Enriquez",
        "criteria": ["El formulario permite crear empresa, dirección, país UE, VAT ID, teléfono y email.", "El VAT ID se valida contra el formato del país seleccionado.", "No se permiten clientes duplicados con el mismo VAT ID.", "El cliente queda activo por defecto.", "Se puede buscar y filtrar por país, estado y nombre de empresa."],
    },
    {
        "key": "RTE-21", "hu": "HU-12", "title": "Gestión de contactos por cliente",
        "actor": "operador", "want": "gestionar múltiples contactos por cliente",
        "purpose": "mantener la comunicación organizada por cada cliente europeo", "assignee": "Johan Steven Muñoz Enriquez",
        "criteria": ["Cada cliente puede tener varios contactos de compras, logística o pagos.", "Cada contacto incluye nombre, cargo, email, teléfono y rol.", "Un contacto puede marcarse como principal.", "Se puede editar y eliminar cada contacto individualmente.", "La ficha del cliente muestra sus contactos."],
    },
    {
        "key": "RTE-24", "hu": "HU-15", "title": "Creación de pedido de exportación",
        "actor": "operador", "want": "crear pedidos con productos, cantidades, precio e Incoterms",
        "purpose": "formalizar las ventas de exportación hacia clientes europeos", "assignee": "stheban hoyos villota",
        "criteria": ["El formulario permite cliente, productos, cantidades, precio unitario, Incoterm y moneda.", "El subtotal y total se calculan automáticamente.", "Se pueden agregar varias líneas de producto a un pedido.", "El pedido se crea en estado Borrador.", "Se valida un contrato activo antes de crear el pedido."],
    },
    {
        "key": "RTE-25", "hu": "HU-16", "title": "Asociación de pedido a lote específico",
        "actor": "operador", "want": "asociar cada pedido a un lote específico",
        "purpose": "garantizar trazabilidad desde origen hasta destino", "assignee": "stheban hoyos villota",
        "criteria": ["Cada línea permite seleccionar un lote disponible.", "Se muestra el peso disponible antes de asignarlo.", "No se permite exceder el peso disponible.", "Al asociarlo, el lote cambia a Reservado.", "Se puede cambiar el lote antes de aprobar.", "La trazabilidad se registra en el historial del lote."],
    },
    {
        "key": "RTE-26", "hu": "HU-17", "title": "Aprobación o rechazo de pedidos",
        "actor": "gerente", "want": "aprobar o rechazar pedidos antes de su confirmación",
        "purpose": "controlar las exportaciones y asegurar que cumplen las políticas de la empresa", "assignee": "stheban hoyos villota",
        "criteria": ["El gerente recibe una notificación al quedar un pedido pendiente.", "Puede aprobar o rechazar con comentarios.", "Al aprobar, el pedido cambia a Confirmado y se reserva el inventario.", "Al rechazar, vuelve a Borrador con el motivo.", "Se registra quién aprobó o rechazó y cuándo.", "El pedido rechazado queda visible al operador con el motivo."],
    },
    {
        "key": "RTE-28", "hu": "HU-19", "title": "Generación de proforma invoice",
        "actor": "operador", "want": "generar una proforma invoice PDF",
        "purpose": "enviarla al cliente como documento preliminar antes de confirmar la venta", "assignee": "stheban hoyos villota",
        "criteria": ["La proforma se genera con los datos del pedido.", "Incluye cliente, productos, cantidades, precios, Incoterms, moneda y fecha de validez.", "Se genera como PDF descargable.", "Tiene un número secuencial único.", "Se puede enviar por email desde el sistema.", "Queda registrada en el historial del pedido."],
    },
    {
        "key": "RTE-30", "hu": "HU-21", "title": "Generación de packing list",
        "actor": "operador", "want": "generar el packing list con el contenido exacto de cada envío",
        "purpose": "cumplir requisitos aduaneros y de documentación de exportación", "assignee": "Supersamuel 909",
        "criteria": ["Se genera automáticamente con los datos del pedido aprobado.", "Incluye descripción, cantidad, pesos bruto/neto y dimensiones.", "Se genera como PDF descargable.", "Se asocia al envío correspondiente.", "Se puede editar antes de generar la versión final.", "La versión final queda en la documentación del envío."],
    },
    {
        "key": "RTE-31", "hu": "HU-22", "title": "Gestión de certificados fitosanitarios",
        "actor": "operador", "want": "gestionar certificados fitosanitarios para cada envío",
        "purpose": "cumplir regulaciones sanitarias de la Unión Europea", "assignee": "Supersamuel 909",
        "criteria": ["Se registra certificado por envío con entidad emisora, número, emisión y vencimiento.", "Se permite adjuntar PDF o imagen.", "Se asocia al envío y al lote.", "Se alerta cuando está próximo a vencer.", "Se lista y busca por rango de fechas y estado."],
    },
    {
        "key": "RTE-32", "hu": "HU-23", "title": "Gestión de certificados de origen",
        "actor": "operador", "want": "gestionar certificados de origen para cada exportación",
        "purpose": "demostrar procedencia del producto y acceder a preferencias arancelarias", "assignee": "Supersamuel 909",
        "criteria": ["Se registra país de origen, tipo de producto, autoridad emisora y número.", "Se asocia al envío y a los productos exportados.", "Se puede adjuntar el certificado PDF.", "Se validan formatos configurados según destino (EUR.1, Form A, etc.).", "Los certificados se historizan para consultas futuras."],
    },
]

ISSUE_BY_KEY = {issue["key"]: issue for issue in ISSUES}

TESTS = [
    ("T10-01", "RTE-10", "Alta de usuario desde formulario con rol Gerente", "Usuario administrador autenticado", "Ingresar nombre, email, rol y contraseña válidos; guardar.", "La cuenta se crea activa, aparece en el directorio y se muestra confirmación.", "La UI creó la cuenta QA temporal como activa y mostró el mensaje de éxito.", "SI"),
    ("T10-02", "RTE-10", "Validar duplicado de email y contraseña débil", "Administrador autenticado; email de prueba existente o nuevo", "Intentar email duplicado y contraseña menor de 8 caracteres/sin mayúscula.", "El servidor rechaza ambos casos y no duplica la cuenta.", "HTTP 409 para duplicado; HTTP 400 para contraseña fuera de política.", "SI"),
    ("T11-01", "RTE-11", "JWT con claims, vencimiento y renovación", "Usuario operador activo", "Iniciar sesión, inspeccionar claims del JWT y solicitar refresh.", "Claims id/email/rol; expiración 24 h; refresh responde con nuevo token.", "Claims y 86.400 s comprobados; endpoint refresh respondió correctamente.", "SI"),
    ("T11-02", "RTE-11", "Error genérico y bloqueo tras cinco intentos", "Usuario QA, nunca cuenta de administración", "Usar clave incorrecta cinco veces; comparar error con email inexistente.", "Error indistinguible y cuenta bloqueada; desbloqueo administrativo permite el acceso.", "HTTP 401 genérico, bloqueo y desbloqueo probados en usuario sintético.", "SI"),
    ("T11-03", "RTE-11", "Logout elimina credenciales locales", "Sesión de administrador abierta en navegador", "Usar Cerrar sesión y comprobar storage del navegador.", "Token y usuario se eliminan y la vista vuelve al login.", "Token y usuario quedaron nulos en localStorage; ruta /login.", "SI"),
    ("T15-01", "RTE-15", "Alta, estado inicial y duplicado de producto", "Operador autenticado", "Crear variedad/origen nuevos y repetir la misma combinación.", "Primer registro activo; duplicado rechazado.", "Alta HTTP 201 y active=true; duplicado HTTP 409.", "SI"),
    ("T15-02", "RTE-15", "Varios certificados asociados al producto", "Producto QA registrado", "Asociar CALIDAD, ORGANICO y FAIR_TRADE; consultar por productId.", "La ficha del producto recupera más de un certificado asociado.", "Tres tipos aparecieron en la consulta de certificados del producto.", "SI"),
    ("T15-03", "RTE-15", "Búsqueda y filtros del catálogo", "Catálogo con productos sintéticos", "Filtrar por texto, familia y estado desde la pantalla.", "La lista solo muestra los productos que coinciden.", "Búsqueda y consulta filtrada por tipo/estado devolvieron el producto esperado.", "SI"),
    ("T16-01", "RTE-16", "Registro de lote y atributos de origen", "Producto registrado", "Registrar código único, peso, cosecha, proceso y ubicación.", "Lote vinculado al producto con campos persistidos.", "Tres lotes QA creados; códigos únicos y campos presentes.", "SI"),
    ("T16-02", "RTE-16", "Análisis de calidad asociado al lote", "Lote QA existente", "Registrar diez atributos de catación y consultar por lote.", "Análisis relacionado al lote con puntuación total calculada.", "Análisis visible por lote; total 80/100.", "SI"),
    ("T16-03", "RTE-16", "Trazabilidad continua y estados hasta destino", "Lote disponible vinculado a pedido", "Recorrer reserva, envío, certificado y ficha de lote.", "La ficha presenta la relación completa desde origen hasta destino.", "Estados probados; LT-DEMO-002 enlaza REQ-DEMO-002, SHP-DEMO-001, Rotterdam Demo y certificados.", "SI"),
    ("T17-01", "RTE-17", "Crear certificado de calidad y adjuntar PDF", "Producto o lote QA", "Crear con número/emisor/fechas; adjuntar y descargar PDF.", "Metadatos y archivo asociado; descarga privada funciona.", "PDF sintético adjuntado y descargado con HTTP 200.", "SI"),
    ("T17-02", "RTE-17", "Alerta visual de vencimiento en 30 días", "Certificado de prueba que vence en 20 días", "Consultar filtro POR_VENCER y revisar estado en la lista.", "El certificado se clasifica como próximo a vencer.", "Consulta expiringInDays=30 devolvió el registro; etiqueta visual definida en UI.", "SI"),
    ("T17-03", "RTE-17", "Eliminar certificado obsoleto con confirmación", "Certificado sintético obsoleto", "Confirmar eliminación y revisar la lista.", "Se elimina solo después de aceptar la confirmación.", "Confirmación UI aceptada; certificado de prueba dejó de aparecer en la API.", "SI"),
    ("T18-01", "RTE-18", "Resumen, estados y porcentaje de disponibilidad", "Base semillada con lotes en varios estados", "Abrir Inventario y comparar peso y disponibilidad por lote.", "El dashboard suma por producto, muestra lotes, estado y porcentaje.", "Inventario agregado cargó; tarjetas muestran kg por estado y disponibilidad.", "SI"),
    ("T18-02", "RTE-18", "Filtros por tipo/estado/fechas y exportación PDF", "Inventario con café, cacao y fechas de cosecha", "Aplicar filtros y usar Imprimir / PDF.", "Listado filtrado y exportable en PDF.", "Filtros combinados aplicados; CSV descargable generado y botón Imprimir llamó la ruta PDF del navegador.", "SI"),
    ("T18-03", "RTE-18", "Actualización al mover un lote", "Pantalla de inventario abierta", "Crear un lote y esperar el ciclo de actualización.", "El nuevo estado y disponibilidad se reflejan automáticamente.", "A los 22 s, la interfaz mostró 400 kg/6 lotes, coincidiendo con la API.", "SI"),
    ("T20-01", "RTE-20", "Alta, VAT por país y unicidad", "Cliente europeo QA", "Crear VAT alemán, repetirlo y consultar el registro activo.", "Formato aceptado/normalizado, repetido rechazado y cliente activo.", "Alta válida, formato DE y duplicado HTTP 409; status ACTIVO.", "SI"),
    ("T20-02", "RTE-20", "Filtros del directorio de clientes", "Clientes de prueba en más de un país", "Filtrar país, estado y texto de compañía.", "Se devuelve solo la coincidencia esperada.", "Consulta combinada por búsqueda, país y ACTIVO devolvió el cliente QA.", "SI"),
    ("T21-01", "RTE-21", "Crear, editar, marcar principal, listar y borrar contacto", "Cliente QA existente", "Crear contacto con todos los campos; editarlo; revisar ficha; borrar.", "Contacto principal visible con cambios y borrado individual.", "Creación, edición, eliminación y visibilidad en ficha comprobadas.", "SI"),
    ("T24-01", "RTE-24", "Pedido de dos líneas con total y estado inicial", "Cliente activo con contrato activo", "Crear pedido con productos, cantidades, precios, moneda e Incoterm.", "Se calcula el total y queda en Borrador.", "Dos líneas guardadas; total EUR 40; estado BORRADOR.", "SI"),
    ("T24-02", "RTE-24", "Pedido sin contrato vigente", "Cliente activo sin contrato activo", "Intentar crear pedido.", "La API rechaza la operación antes de crear el pedido.", "HTTP 400 con requisito de contrato activo.", "SI"),
    ("T25-01", "RTE-25", "Asignación de lote y límite de peso", "Lote disponible del mismo producto", "Asociar lote y probar una cantidad superior a su peso.", "El lote queda reservado; el exceso se rechaza.", "Reserva confirmada; pedido que excedía el peso recibió HTTP 400.", "SI"),
    ("T25-02", "RTE-25", "Cambiar lote antes de aprobar", "Pedido borrador con lote A y lote B disponible", "Reasignar la línea de A a B.", "A vuelve a Disponible y B queda Reservado.", "Estados DISPONIBLE/RESERVADO comprobados.", "SI"),
    ("T25-03", "RTE-25", "Historial de cambios de lote", "Pedido con reasignación", "Consultar la ficha histórica del lote anterior y nuevo.", "El historial conserva la asociación anterior además de la actual.", "La reasignación registró eventos de reserva/liberación y aparecen en los pasaportes de ambos lotes.", "SI"),
    ("T26-01", "RTE-26", "Notificar, aprobar y guardar responsable/fecha", "Pedido Borrador y rol gerente/admin", "Enviar a aprobación; revisar notificación; aprobar.", "Notificación, estado CONFIRMADO, reserva y actor/fecha registrados.", "Notificación presente; status, approvedById y approvedAt guardados.", "SI"),
    ("T26-02", "RTE-26", "Rechazar y devolver a borrador", "Pedido pendiente con lote reservado", "Rechazar con comentario; revisar desde operador.", "Vuelve a Borrador con motivo visible; lote se libera.", "BORRADOR, rejectReason y lote DISPONIBLE comprobados por API.", "SI"),
    ("T28-01", "RTE-28", "Proforma PDF y contenido comercial", "Pedido QA y cliente asociado", "Generar PDF con validez futura.", "PDF incluye cliente, líneas, precios, Incoterm, moneda y validez.", "PDF descargable de 2.162 bytes; contenido cargado desde pedido; revisión textual pendiente.", "SI"),
    ("T28-02", "RTE-28", "Número y registro histórico de proforma", "Pedido sin o con proforma previa", "Generar consecutivo y consultar facturas del pedido.", "Número único registrado en el historial.", "PI-2026-00002 registrado en Invoice e historial.", "SI"),
    ("T28-03", "RTE-28", "Enviar proforma por correo", "SMTP y correo QA configurados", "Enviar desde el expediente y comprobar recepción.", "Email entregado con PDF adjunto y estado visible.", "No ejecutado: requiere SMTP aprobado; no se envió correo externo.", "PENDIENTE"),
    ("T30-01", "RTE-30", "Editar embalaje antes del PDF", "Pedido editable con línea asignada", "Editar bultos/peso bruto/dimensiones y generar packing list.", "Datos válidos incluidos; peso bruto inferior al neto rechazado.", "Edición y validación neto/bruto probadas.", "SI"),
    ("T30-02", "RTE-30", "Generar packing list solo desde pedido aprobado y asociado a envío", "Pedido aprobado con datos completos y ShipmentOrder", "Intentar en borrador/sin envío; después generar en pedido confirmado con envío.", "Los casos inválidos se bloquean; el válido entrega PDF vinculado al envío.", "HTTP 400 en borrador y aprobado sin envío; PDF de pedido confirmado con envío, HTTP 200.", "SI"),
    ("T30-03", "RTE-30", "Versionado e historial documental", "Pedido confirmado asociado a envío", "Generar dos veces y consultar documentos del pedido.", "Ambas versiones tienen números únicos y vínculo al envío.", "PL-REQ-DEMO-002-01 y -02 registrados; enlace SHP-DEMO-001.", "SI"),
    ("T31-01", "RTE-31", "Certificado fitosanitario enlazado y fechado", "Lote y envío existentes", "Registrar emisor, número, fechas y asociar lote/envío.", "El registro queda asociado al envío y al lote.", "Certificado QA vinculado a ambos; fechas persistidas.", "SI"),
    ("T31-02", "RTE-31", "Adjuntar/descargar archivo fitosanitario", "Certificado QA creado", "Subir PDF sintético y descargarlo.", "El adjunto puede recuperarse con el tipo de contenido correcto.", "PDF guardado y descarga HTTP 200.", "SI"),
    ("T31-03", "RTE-31", "Buscar por vencimiento y rango de fechas", "Certificados con distintas fechas", "Consultar estado POR_VENCER y rango de emisión.", "La lista respeta ambos filtros.", "Consulta combinada retornó registro próximo a vencer.", "SI"),
    ("T32-01", "RTE-32", "Validar formatos de origen por destino", "Envío a destino con formatos configurados", "Intentar formato inválido y luego formato permitido.", "El formato no permitido se rechaza y el permitido se guarda.", "INVALIDO devolvió HTTP 400; EUR.1 creó certificado.", "SI"),
    ("T32-02", "RTE-32", "Asociar origen, producto/lote, envío, adjunto e historial", "Producto, lote y envío QA", "Registrar certificado, asociar y subir PDF; consultar historial.", "El expediente conserva asociaciones y archivo descargable.", "Certificado asociado a producto/lote y envío; PDF sintético descargado.", "SI"),
]


def font(size: int):
    return ImageFont.truetype(FONT_PATH, size)


def wrapped(draw, text, width, size):
    f = font(size)
    words, lines, current = text.split(), [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=f)[2] <= width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_case_diagram(path: Path):
    w, h = 1900, 1340
    im = Image.new("RGB", (w, h), "#F4F3EF")
    d = ImageDraw.Draw(im)
    d.text((76, 45), "UML · CASOS DE USO", font=font(34), fill="#102B46")
    d.text((78, 96), "Flujo interno de operación exportadora · historias RTE del Sprint 2", font=font(20), fill="#667085")
    boundary = (350, 160, 1550, 1260)
    d.rounded_rectangle(boundary, radius=24, outline="#8DA0B5", width=3, fill="#FCFCFA")
    d.text((385, 178), "RiTech Export Desk", font=font(22), fill="#3465E9")
    nodes = [
        ("Registrar usuarios", 420, 240), ("Mantener catálogo y lotes", 800, 240), ("Consultar inventario", 1180, 240),
        ("Administrar clientes", 420, 425), ("Gestionar contactos/contratos", 800, 425), ("Crear pedido", 1180, 425),
        ("Asignar lote y embalaje", 420, 610), ("Aprobar / rechazar pedido", 800, 610), ("Generar proforma", 1180, 610),
        ("Gestionar certificados", 420, 795), ("Generar packing list", 800, 795), ("Vincular envío", 1180, 795),
        ("Consultar trazabilidad", 420, 980), ("Exportar inventario", 800, 980), ("Descargar documentos", 1180, 980),
    ]
    # Actor-to-system connectors sit behind the use-case ovals.
    actor_xy = {"Administrador": (155, 250), "Operador": (155, 455), "Gerente": (155, 675), "Contabilidad": (155, 900), "Logística / cliente": (1710, 1010)}
    actor_links = {
        "Administrador": [0, 2, 7], "Operador": [1, 3, 4, 5, 6, 9, 10], "Gerente": [2, 7, 12],
        "Contabilidad": [8, 14], "Logística / cliente": [11, 14],
    }
    for actor, indices in actor_links.items():
        sx, sy = actor_xy[actor]
        for index in indices:
            _, x, y = nodes[index]
            target_x = x if actor != "Logística / cliente" else x + 300
            start_x = sx + (45 if actor != "Logística / cliente" else -45)
            d.line((start_x, sy + 38, target_x, y + 47), fill="#A8B5C3", width=2)
    for index, (label, x, y) in enumerate(nodes):
        d.rounded_rectangle((x, y, x + 300, y + 96), radius=48, fill="#E8F3EF", outline="#6E968B", width=2)
        lines = wrapped(d, label, 258, 17)
        text_y = y + (96 - len(lines) * 23) / 2
        for line in lines:
            box = d.textbbox((0, 0), line, font=font(17))
            d.text((x + (300 - (box[2] - box[0])) / 2, text_y), line, font=font(17), fill="#19334C")
            text_y += 23
    # Simple actor pictograms and role labels.
    for actor, (x, y) in actor_xy.items():
        d.ellipse((x - 18, y - 28, x + 18, y + 8), outline="#19334C", width=3)
        d.line((x, y + 8, x, y + 63), fill="#19334C", width=3)
        d.line((x - 28, y + 26, x + 28, y + 26), fill="#19334C", width=3)
        d.line((x, y + 63, x - 24, y + 94), fill="#19334C", width=3)
        d.line((x, y + 63, x + 24, y + 94), fill="#19334C", width=3)
        text_w = d.textbbox((0, 0), actor, font=font(14))[2]
        d.text((x - text_w / 2, y + 103), actor, font=font(14), fill="#19334C")
    im.save(path)


def draw_data_model(path: Path):
    w, h = 1900, 1290
    im = Image.new("RGB", (w, h), "#F4F3EF")
    d = ImageDraw.Draw(im)
    d.text((70, 40), "MODELO DE DATOS · RI TECH SPRINT 2", font=font(32), fill="#102B46")
    d.text((72, 88), "Entidades relacionales principales · PK / FK resumidas a partir del esquema Prisma", font=font(19), fill="#667085")
    cols = [75, 525, 975, 1425]
    rows = [160, 425, 690, 955]
    box_w, box_h = 365, 205
    entities = [
        ("User", ["id · PK", "email · UNIQUE", "roles[] · active", "approvedOrders / notifications"]),
        ("Product", ["id · PK", "type · variety · origin", "UNIQUE(variety, origin)", "lots · certificates"]),
        ("Lot", ["id · PK · traceabilityCode", "productId · FK", "weight · dates · origin", "status · certificates · items"]),
        ("Client", ["id · PK · userId FK?", "company · country · vatId", "contacts · contracts", "orders"]),
        ("ClientContact", ["id · PK · clientId FK", "name · position · email", "phone · role · isPrimary"]),
        ("ClientContract", ["id · PK · clientId FK", "contractNumber · UNIQUE", "status · start/end dates"]),
        ("Order", ["id · PK · orderNumber", "clientId · FK · status", "currency · incoterm · total", "items · invoices · documents"]),
        ("OrderItem", ["id · PK · orderId FK", "productId FK · lotId FK?", "quantity · unitPrice", "packageCount · grossWeight · dimensions"]),
        ("Invoice", ["id · PK · orderId FK", "invoiceNumber · UNIQUE", "type · status · amount", "validUntil"]),
        ("Certificate", ["id · PK · type · number", "productId / lotId / shipmentId", "issuer · dates · format", "fileUrl · notes"]),
        ("Shipment", ["id · PK · shipmentNumber", "destinationId · FK", "status · container · dates", "orders · documents · certificates"]),
        ("ShipmentOrder", ["id · PK", "shipmentId · FK", "orderId · FK", "UNIQUE(shipmentId, orderId)"]),
        ("Document", ["id · PK · orderId FK", "shipmentId · FK?", "type · number · status", "fileUrl · metadata · dates"]),
        ("ShippingDestination", ["id · PK", "city · country · portCode", "certificateFormats[]", "default mode / Incoterm"]),
        ("QualityAnalysis", ["id · PK · lotId FK", "analyst · analyzedAt", "10 SCA score fields", "totalScore · defects"]),
    ]
    positions = {name: (cols[i % 4], rows[i // 4]) for i, (name, _) in enumerate(entities)}
    links = [
        ("User", "Order", "seller / approval"), ("Client", "ClientContact", "1:N"), ("Client", "ClientContract", "1:N"),
        ("Client", "Order", "1:N"), ("Product", "Lot", "1:N"), ("Product", "OrderItem", "1:N"),
        ("Lot", "OrderItem", "0:N"), ("Order", "OrderItem", "1:N"), ("Order", "Invoice", "1:N"),
        ("Order", "Document", "1:N"), ("Order", "ShipmentOrder", "1:N"), ("Shipment", "ShipmentOrder", "1:N"),
        ("ShippingDestination", "Shipment", "1:N"), ("Shipment", "Certificate", "1:N"),
        ("Lot", "Certificate", "1:N"), ("Lot", "QualityAnalysis", "1:N"), ("Shipment", "Document", "1:N"),
    ]
    for left, right, label in links:
        x1, y1 = positions[left]
        x2, y2 = positions[right]
        start = (x1 + box_w, y1 + box_h / 2) if x2 > x1 else (x1 + box_w / 2, y1 + box_h)
        end = (x2, y2 + box_h / 2) if x2 > x1 else (x2 + box_w / 2, y2)
        d.line((start[0], start[1], end[0], end[1]), fill="#BCC6CF", width=2)
    for name, fields in entities:
        x, y = positions[name]
        d.rounded_rectangle((x, y, x + box_w, y + box_h), radius=14, fill="#FFFFFF", outline="#A8B5C3", width=2)
        d.rounded_rectangle((x, y, x + box_w, y + 47), radius=14, fill="#102B46")
        d.rectangle((x, y + 30, x + box_w, y + 47), fill="#102B46")
        d.text((x + 16, y + 12), name, font=font(19), fill="#FFFFFF")
        ty = y + 60
        for field in fields:
            d.text((x + 16, ty), field, font=font(14), fill="#19334C")
            ty += 32
    im.save(path)


def add_docx_header(doc: Document, label: str):
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    header = section.header.paragraphs[0]
    header.text = f"RITECH  /  CALIDAD DE SOFTWARE                                      {label.upper()}"
    header.style = doc.styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
    header.runs[0].font.bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("RiTech · Sprint 2 · 27/09/2026  |  ")
    field = footer.add_run()
    field._r.append(__import__("docx").oxml.OxmlElement("w:fldSimple"))


def style_docx(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for style_name, size, color in [("Title", 30, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 10, ORANGE)]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = __import__("docx").oxml.OxmlElement("w:shd")
    shd.set(__import__("docx").oxml.ns.qn("w:fill"), fill)
    tc_pr.append(shd)


def set_docx_cell(cell, text, bold=False, color=INK, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_docx_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(62)
    r = p.add_run("RiTech Export Desk")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(ORANGE)
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(subtitle)
    p.style = doc.styles["Subtitle"]
    p.runs[0].font.size = Pt(15)
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_paragraph(f"Entrega: 30 de septiembre de 2026, 2:00 p. m.  ·  Corte: {DATE_LABEL}")
    doc.add_paragraph("Documento de trabajo · datos de demostración sintéticos")
    doc.add_page_break()


def make_requirements_docx():
    doc = Document()
    style_docx(doc)
    add_docx_header(doc, "Especificación de requisitos")
    add_docx_title(doc, "Especificación de requisitos", "15 historias seleccionadas · Sprint 2")
    doc.add_heading("1. Propósito y contexto", 1)
    doc.add_paragraph("RiTech Export Desk organiza la operación interna para exportar café y cacao colombianos a compradores europeos. El alcance de este documento sigue las historias seleccionadas en el Sprint 2 del tablero RTE y describe el prototipo local, sus criterios de aceptación, actores, modelo de datos y evidencia de pruebas.")
    doc.add_heading("2. Alcance", 1)
    doc.add_paragraph("Se incluyen RTE-10, RTE-11, RTE-15, RTE-16, RTE-17, RTE-18, RTE-20, RTE-21, RTE-24, RTE-25, RTE-26, RTE-28, RTE-30, RTE-31 y RTE-32. Las seis historias retiradas por complejidad y no incluidas en este corte son RTE-12, RTE-33, RTE-34, RTE-38, RTE-40 y RTE-44.")
    doc.add_paragraph("El contrato activo se representa en el prototipo como una precondición de RTE-24. Esto no equivale a afirmar que la historia completa de contratos marco esté terminada.")
    doc.add_heading("3. Actores", 1)
    actor_table = doc.add_table(rows=1, cols=2)
    actor_table.style = "Light Shading Accent 1"
    for i, text in enumerate(["Actor", "Responsabilidad dentro del alcance"]):
        set_docx_cell(actor_table.rows[0].cells[i], text, True, "FFFFFF", 9)
        shade_cell(actor_table.rows[0].cells[i], NAVY)
    for actor, role in [
        ("Administrador", "Gestiona usuarios y acceso; también puede revisar pedidos."),
        ("Operador", "Registra catálogo, lotes, clientes, pedidos y documentos."),
        ("Gerente", "Consulta existencias y aprueba o rechaza pedidos."),
        ("Contabilidad", "Consulta/genera documentación comercial y proformas."),
        ("Logística", "Asocia envíos y gestiona documentación de despacho."),
        ("Comprador europeo", "Consulta documentos vinculados cuando su rol lo permite."),
    ]:
        cells = actor_table.add_row().cells
        set_docx_cell(cells[0], actor, True)
        set_docx_cell(cells[1], role)
    doc.add_page_break()
    doc.add_heading("4. Diagrama UML de casos de uso", 1)
    doc.add_paragraph("Diagrama de alcance para los actores y funciones principales. Los casos se agrupan en el flujo operativo interno.")
    doc.add_picture(str(DIAGRAMS / "casos-de-uso.png"), width=Inches(6.7))
    doc.add_heading("5. Requisitos funcionales e historias de usuario", 1)
    for issue in ISSUES:
        doc.add_heading(f"{issue['key']} · {issue['title']}", 2)
        info = doc.add_table(rows=0, cols=2)
        info.style = "Light Shading Accent 1"
        for label, value in [("Como", issue["actor"]), ("Quiero", issue["want"]), ("Para", issue["purpose"]), ("Asignado en Jira (corte)", issue["assignee"])]:
            row = info.add_row().cells
            set_docx_cell(row[0], label, True, NAVY)
            set_docx_cell(row[1], value)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run("Criterios de aceptación").bold = True
        for criterion in issue["criteria"]:
            doc.add_paragraph(criterion, style="List Bullet")
        source = doc.add_paragraph(f"Fuente de alcance: https://jhonra2008.atlassian.net/browse/{issue['key']}")
        source.style = doc.styles["Caption"]
    doc.add_page_break()
    doc.add_heading("6. Modelo de datos", 1)
    doc.add_paragraph("Resumen de entidades y relaciones principales del esquema Prisma/PostgreSQL. Las llaves y cardinalidades representadas resumen el modelo; el esquema del repositorio es la fuente técnica de detalle.")
    doc.add_picture(str(DIAGRAMS / "modelo-de-datos.png"), width=Inches(6.7))
    doc.add_heading("Entidades centrales", 2)
    doc.add_paragraph("User administra roles; Product agrupa lotes y certificados; Client mantiene contactos y contratos; Order contiene líneas de pedido, facturas y documentos; Shipment conecta pedidos con destino, documentos y certificados; QualityAnalysis conserva atributos de calidad por lote.")
    doc.add_paragraph("La base local de demostración es `ritech_sprint2`; el esquema fuente es `apps/api/prisma/schema.prisma`. La carga seed contiene registros sintéticos. La base anterior `ritech_db` se conserva.")
    doc.add_heading("7. Requisitos no funcionales propuestos (por validar)", 1)
    for item in [
        "Seguridad: roles controlan las operaciones; contraseñas almacenadas con hash; JWT de 24 horas con renovación.",
        "Integridad: VAT y trazabilidad tienen restricciones de unicidad; los pedidos requieren contrato activo y no pueden superar el peso del lote.",
        "Trazabilidad: lote, línea de pedido, certificado, envío y documentos se enlazan mediante identificadores del modelo.",
        "Usabilidad: interfaz adaptable a pantallas estrechas, con estados, validaciones y formato de impresión para inventario/documentos.",
        "Operación: el prototipo usa PostgreSQL local y almacenamiento de archivos local; el despliegue real debe definir persistencia, respaldo y configuración de SMTP.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("8. Supuestos y límites", 1)
    doc.add_paragraph("Los datos seed y los adjuntos sintéticos son para pruebas, no son documentos comerciales. La validación VAT verifica estructura/prefijo configurados; no consulta VIES. Los formatos de origen se controlan por una lista del destino y requieren validación operativa/comercial antes de su uso real. El envío de email necesita SMTP configurado y autorizado.")
    doc.add_paragraph("Los resultados parciales y pendientes se detallan en `Plan-de-Pruebas-Sprint-2.xlsx`. El documento separado de aportes marca hipótesis basadas en asignaciones Jira y no las presenta como evidencia de trabajo individual.")
    doc.save(OUT / "Requerimientos-Sprint-2.docx")


def make_contributions_docx():
    doc = Document()
    style_docx(doc)
    add_docx_header(doc, "Supuestos de aportes")
    add_docx_title(doc, "Aportes del equipo", "Distribución supuesta · documento separado del historial de trabajo real")
    doc.add_heading("Propósito y criterio de atribución", 1)
    doc.add_paragraph("Este documento presenta una propuesta de distribución para preparar la entrega del Sprint 2. Se elaboró con las asignaciones visibles en Jira al 27/09/2026 y el tema de las historias. La asignación no demuestra que una persona haya ejecutado, revisado o terminado el trabajo durante septiembre.")
    doc.add_paragraph("Durante esta revisión se actualizaron 14 de las 15 historias seleccionadas a Finalizada con comentarios de verificación y se dejó RTE-28 En curso por el requisito de email/SMTP. Esas transiciones registran el estado verificado en esta revisión, pero no prueban aportes individuales durante septiembre. Cada integrante debe confirmar, corregir o retirar su fila antes de la presentación.")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    for i, value in enumerate(["Integrante asignado en Jira", "Historias del alcance", "Aporte propuesto (supuesto)", "Validación del integrante"]):
        set_docx_cell(table.rows[0].cells[i], value, True, "FFFFFF", 8)
        shade_cell(table.rows[0].cells[i], NAVY)
    rows = [
        ("JhonRamirez22", "RTE-10, RTE-11", "Propuesta: revisar registro de usuarios, perfiles, autenticación, sesión y evidencias de seguridad.", "Pendiente de confirmar"),
        ("Miguel Felipe Ceballos Ramirez", "RTE-15, RTE-16, RTE-17, RTE-18", "Propuesta: revisar catálogo, trazabilidad de lotes, certificados de calidad e inventario.", "Pendiente de confirmar"),
        ("Johan Steven Muñoz Enriquez", "RTE-20, RTE-21", "Propuesta: revisar registro de clientes europeos, validación VAT y contactos.", "Pendiente de confirmar"),
        ("stheban hoyos villota", "RTE-24, RTE-25, RTE-26, RTE-28", "Propuesta: revisar pedidos, vínculo a lote, aprobación y documentación proforma.", "Pendiente de confirmar"),
        ("Supersamuel 909", "RTE-30, RTE-31, RTE-32", "Propuesta: revisar packing list y certificados fitosanitarios/de origen.", "Pendiente de confirmar"),
    ]
    for data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(data):
            set_docx_cell(cells[i], value, bold=(i == 0), size=8)
    doc.add_heading("Evidencia que falta para afirmar aportes reales", 1)
    doc.add_paragraph("Antes de reemplazar los supuestos por un resumen real, acordar con el equipo la confirmación individual y recopilar enlaces a commits/PR, comentarios técnicos, revisiones, ejecuciones de pruebas, decisiones y fechas. Mantener por separado el trabajo asignado, el trabajo ejecutado y los resultados verificables.")
    doc.add_heading("Fuente", 1)
    doc.add_paragraph("Asignaciones y estados consultados en los issues RTE-10, RTE-11, RTE-15, RTE-16, RTE-17, RTE-18, RTE-20, RTE-21, RTE-24, RTE-25, RTE-26, RTE-28, RTE-30, RTE-31 y RTE-32. Corte: 27 de septiembre de 2026.")
    doc.save(OUT / "Supuesto-Trabajo-Integrantes-Sprint-2.docx")


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=26, leading=31, textColor=colors.HexColor("#" + NAVY), spaceAfter=12))
    styles.add(ParagraphStyle(name="H1x", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=colors.HexColor("#" + NAVY), spaceBefore=12, spaceAfter=8))
    styles.add(ParagraphStyle(name="H2x", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#" + BLUE), spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="Bodyx", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.6, leading=12, spaceAfter=5, textColor=colors.HexColor("#" + INK)))
    styles.add(ParagraphStyle(name="Smallx", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.2, leading=9, textColor=colors.HexColor("#" + GRAY)))
    styles.add(ParagraphStyle(name="Cellx", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.4, leading=9, textColor=colors.HexColor("#" + INK)))
    styles.add(ParagraphStyle(name="CellWhite", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=colors.white))
    return styles


def para(text, style, markup=False):
    content = str(text) if markup else escape(str(text))
    return Paragraph(content.replace("\n", "<br/>"), style)


def on_page(canvas, doc):
    canvas.saveState()
    if doc.page > 1:
        canvas.setStrokeColor(colors.HexColor("#D6DEE6"))
        canvas.line(18 * mm, 17 * mm, 192 * mm, 17 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#667085"))
        canvas.drawString(18 * mm, 11 * mm, "RiTech · Sprint 2 · 27/09/2026")
        canvas.drawRightString(192 * mm, 11 * mm, f"Página {doc.page}")
    canvas.restoreState()


def pdf_table(rows, widths, header=True):
    styles = pdf_styles()
    data = [[para(cell, styles["CellWhite"] if header and row_index == 0 else styles["Cellx"]) for cell in row] for row_index, row in enumerate(rows)]
    table = Table(data, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if header:
        commands.extend([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)])
        for row in range(1, len(rows)):
            if row % 2 == 0:
                commands.append(("BACKGROUND", (0, row), (-1, row), colors.HexColor("#F4F3EF")))
    table.setStyle(TableStyle(commands))
    return table


def make_requirements_pdf():
    styles = pdf_styles()
    story = [Spacer(1, 42 * mm), para("RITECH EXPORT DESK", styles["H2x"]), Spacer(1, 4 * mm), para("Especificación de requisitos", styles["CoverTitle"]), para("15 historias seleccionadas · Sprint 2", styles["Heading2"]), Spacer(1, 10 * mm), para(f"Entrega: 30 de septiembre de 2026, 2:00 p. m.\nCorte: {DATE_LABEL}\nDocumento de trabajo · datos sintéticos", styles["Bodyx"]), PageBreak()]
    story += [para("1. Propósito y contexto", styles["H1x"]), para("RiTech Export Desk organiza la operación interna para exportar café y cacao colombianos a compradores europeos. Este documento sigue las historias seleccionadas en Sprint 2 y describe actores, criterios, modelo de datos, supuestos y evidencia de pruebas.", styles["Bodyx"]), para("2. Alcance", styles["H1x"]), para("Incluye RTE-10, RTE-11, RTE-15, RTE-16, RTE-17, RTE-18, RTE-20, RTE-21, RTE-24, RTE-25, RTE-26, RTE-28, RTE-30, RTE-31 y RTE-32. RTE-12, RTE-33, RTE-34, RTE-38, RTE-40 y RTE-44 quedaron fuera del alcance seleccionado por complejidad. El contrato activo se usa como precondición técnica de RTE-24; esto no significa que la historia completa de contratos marco esté terminada.", styles["Bodyx"]), para("3. Actores", styles["H1x"]), pdf_table([["Actor", "Responsabilidad"]] + [[a, b] for a, b in [("Administrador", "Usuarios, roles y revisión de pedidos."), ("Operador", "Catálogo, lotes, clientes, pedidos y documentos."), ("Gerente", "Inventario y decisión de aprobar/rechazar."), ("Contabilidad", "Documentación comercial y proformas."), ("Logística", "Envíos y documentos de despacho."), ("Comprador UE", "Consulta de documentos cuando su rol lo permite.")]], [43 * mm, 132 * mm]), PageBreak()]
    story += [para("4. Diagrama UML de casos de uso", styles["H1x"]), para("Resumen del flujo y actores del alcance seleccionado.", styles["Bodyx"]), PDFImage(str(DIAGRAMS / "casos-de-uso.png"), width=175 * mm, height=123 * mm), PageBreak(), para("5. Requisitos funcionales", styles["H1x"])]
    for issue in ISSUES:
        blocks = [para(f"{issue['key']} · {issue['title']}", styles["H2x"]), para(f"<b>Como:</b> {escape(issue['actor'])}<br/><b>Quiero:</b> {escape(issue['want'])}<br/><b>Para:</b> {escape(issue['purpose'])}<br/><b>Asignación Jira (corte):</b> {escape(issue['assignee'])}", styles["Bodyx"], markup=True)]
        blocks.extend(para("• " + c, styles["Bodyx"]) for c in issue["criteria"])
        blocks.append(para(f"Fuente Jira: {issue['key']} · https://jhonra2008.atlassian.net/browse/{issue['key']}", styles["Smallx"]))
        story.append(KeepTogether(blocks))
    story += [PageBreak(), para("6. Modelo de datos", styles["H1x"]), para("Resumen de entidades y relaciones principales del esquema Prisma/PostgreSQL; el esquema del repositorio conserva el detalle técnico.", styles["Bodyx"]), PDFImage(str(DIAGRAMS / "modelo-de-datos.png"), width=175 * mm, height=119 * mm), para("Entidades centrales: User, Product, Lot, Client, ClientContact, ClientContract, Order, OrderItem, Invoice, Certificate, Shipment, ShipmentOrder, Document, ShippingDestination y QualityAnalysis.", styles["Bodyx"]), para("La base de demostración es ritech_sprint2. Su fuente de esquema es apps/api/prisma/schema.prisma. La base ritech_db se conserva. Seed y archivos de prueba son sintéticos.", styles["Bodyx"]), para("7. Requisitos no funcionales propuestos (por validar)", styles["H1x"])]
    for item in ["Seguridad: roles, contraseñas con hash y JWT con renovación.", "Integridad: VAT/trazabilidad únicos, contrato activo para pedido y límite de peso por lote.", "Trazabilidad: relaciones entre lote, líneas, certificados, envío y documentos.", "Usabilidad: diseño adaptable y controles para filtrar/imprimir.", "Operación: despliegue real requiere almacenamiento persistente, respaldo y SMTP configurado."]:
        story.append(para("• " + item, styles["Bodyx"]))
    story += [para("8. Supuestos y límites", styles["H1x"]), para("Los registros y adjuntos del seed son de demostración, no comerciales. La validación VAT revisa formato/prefijo configurado y no consulta VIES. Los formatos de origen dependen de la lista configurada por destino y requieren validación operativa antes de un uso real. El envío de correo requiere SMTP; no se envió correo externo durante las pruebas.", styles["Bodyx"]), para("Los resultados por caso están en Plan-de-Pruebas-Sprint-2.xlsx. Las hipótesis de aportes individuales están en un archivo separado y requieren confirmación del equipo.", styles["Bodyx"])]
    SimpleDocTemplate(str(OUT / "Requerimientos-Sprint-2.pdf"), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=17*mm, bottomMargin=23*mm, title="Requerimientos Sprint 2 RiTech").build(story, onFirstPage=on_page, onLaterPages=on_page)


def make_contributions_pdf():
    styles = pdf_styles()
    story = [Spacer(1, 45*mm), para("RITECH EXPORT DESK", styles["H2x"]), Spacer(1, 4*mm), para("Aportes del equipo", styles["CoverTitle"]), para("Distribución supuesta · documento separado del historial real", styles["Heading2"]), Spacer(1, 10*mm), para(f"Corte: {DATE_LABEL} · Entrega: 30/09/2026, 2:00 p. m.", styles["Bodyx"]), PageBreak(), para("Criterio de atribución", styles["H1x"]), para("Propuesta de distribución elaborada desde asignaciones visibles en Jira y el tema de cada historia. La asignación no demuestra ejecución, revisión ni finalización individual. En esta revisión se actualizaron 14 historias a Finalizada con comentarios de verificación y RTE-28 quedó En curso por SMTP; esas transiciones no identifican quién implementó o probó cada criterio. Cada integrante debe confirmar o corregir su fila antes de presentarla.", styles["Bodyx"])]
    rows = [["Integrante Jira", "Historias", "Aporte propuesto (supuesto)", "Validación"]]
    rows += [
        ["JhonRamirez22", "RTE-10, RTE-11", "Revisar alta de usuarios, roles, autenticación y sesión.", "Pendiente"],
        ["Miguel Felipe Ceballos Ramirez", "RTE-15, RTE-16, RTE-17, RTE-18", "Revisar catálogo, lotes, certificados de calidad e inventario.", "Pendiente"],
        ["Johan Steven Muñoz Enriquez", "RTE-20, RTE-21", "Revisar clientes europeos, VAT y contactos.", "Pendiente"],
        ["stheban hoyos villota", "RTE-24, RTE-25, RTE-26, RTE-28", "Revisar pedidos, lote, aprobación y proforma.", "Pendiente"],
        ["Supersamuel 909", "RTE-30, RTE-31, RTE-32", "Revisar packing list y certificados de exportación.", "Pendiente"],
    ]
    story += [pdf_table(rows, [36*mm, 33*mm, 75*mm, 31*mm]), Spacer(1, 8*mm), para("Para afirmar aportes reales", styles["H1x"]), para("Solicitar confirmación individual y enlazar commits/PR, revisiones, comentarios técnicos, decisiones y ejecuciones de pruebas. Mantener separados el trabajo asignado, el ejecutado y los resultados comprobables.", styles["Bodyx"]), para("Fuente: asignaciones consultadas en RTE-10, RTE-11, RTE-15, RTE-16, RTE-17, RTE-18, RTE-20, RTE-21, RTE-24, RTE-25, RTE-26, RTE-28, RTE-30, RTE-31 y RTE-32, corte 27/09/2026.", styles["Smallx"])]
    SimpleDocTemplate(str(OUT / "Supuesto-Trabajo-Integrantes-Sprint-2.pdf"), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=17*mm, bottomMargin=23*mm, title="Supuesto de aportes RiTech Sprint 2").build(story, onFirstPage=on_page, onLaterPages=on_page)


def fill_test_workbook():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"No se encontró la plantilla original: {TEMPLATE}")
    wb = load_workbook(TEMPLATE)
    ws = wb["Hoja1"] if "Hoja1" in wb.sheetnames else wb.active
    # Preserve the supplied title and layout; populate the existing team line and test table.
    member_cell = next((c for c in ws[4] if c.value and "integrantes" in str(c.value).lower()), None)
    if member_cell:
        target_col = member_cell.column + 1
        while target_col <= ws.max_column and ws.cell(4, target_col).value:
            target_col += 1
        ws.cell(4, target_col, "JhonRamirez22 · Miguel Felipe Ceballos Ramirez · Johan Steven Muñoz Enriquez · stheban hoyos villota · Supersamuel 909")
        ws.cell(4, target_col).alignment = Alignment(wrap_text=True, vertical="center")
    headers = ["ID", "Historia", "Tipo de prueba", "Escenario / caso", "Precondiciones", "Pasos", "Resultado esperado", "Resultado obtenido", "Responsable Jira", "Fecha de ejecución", "Pasa (SI/NO/PENDIENTE)", "Observaciones"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(6, col, value)
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    # Clear example rows while retaining their formatting.
    for row in range(7, max(ws.max_row, 7 + len(TESTS)) + 1):
        for col in range(1, 13):
            cell = ws.cell(row, col)
            if row >= 7:
                cell.value = None
    widths = [12, 12, 21, 38, 31, 38, 41, 41, 30, 18, 24, 29]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[__import__("openpyxl").utils.get_column_letter(col)].width = width
    for idx, test in enumerate(TESTS, 7):
        test_id, issue_key, scenario, pre, steps, expected, actual, status = test
        issue = ISSUE_BY_KEY[issue_key]
        values = [test_id, issue_key, "Integración / aceptación", scenario, pre, steps, expected, actual, issue["assignee"], date(2026, 9, 27) if status == "SI" else None, status, "Prueba API/flujo local; todos los datos son sintéticos." if status == "SI" else "Pendiente: no equivale a fallo."]
        for col, value in enumerate(values, 1):
            cell = ws.cell(idx, col, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name="Arial", size=8, color=INK)
            cell.border = Border(bottom=Side(style="hair", color="D6DEE6"))
            if col == 11:
                cell.font = Font(name="Arial", size=8, bold=True, color="18794E" if status == "SI" else "9A6700")
                cell.fill = PatternFill("solid", fgColor="E8F3EF" if status == "SI" else "FFF4D6")
        ws.row_dimensions[idx].height = 64
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A6:L{6 + len(TESTS)}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "1:6"
    wb.save(OUT / "Plan-de-Pruebas-Sprint-2.xlsx")


def make_synthetic_certificate():
    from reportlab.lib.pagesizes import A4 as PAGE
    from reportlab.pdfgen import canvas
    path = FIXTURES / "certificado-sintetico-demo.pdf"
    c = canvas.Canvas(str(path), pagesize=PAGE)
    width, height = PAGE
    c.setFillColor(colors.HexColor("#" + NAVY))
    c.rect(0, height - 42 * mm, width, 42 * mm, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 17)
    c.drawString(20 * mm, height - 24 * mm, "RiTech · ADJUNTO SINTÉTICO")
    c.setFillColor(colors.HexColor("#" + ORANGE))
    c.setFont("Helvetica-Bold", 12)
    c.drawString(20 * mm, height - 58 * mm, "MUESTRA PARA PRUEBAS · NO ES UN CERTIFICADO OFICIAL")
    c.setFillColor(colors.HexColor("#" + INK))
    c.setFont("Helvetica", 10)
    lines = ["Identificador: SAMPLE-QA-01", "Producto: Café de muestra (registro ficticio)", "Lote: SAMPLE-LOT-01", "Emisor: Entidad de prueba ficticia", "Fecha: 27/09/2026", "", "Este archivo solo valida carga, descarga y asociación de adjuntos.", "No tiene validez comercial, sanitaria, aduanera ni legal."]
    y = height - 76 * mm
    for line in lines:
        c.drawString(20 * mm, y, line)
        y -= 9 * mm
    c.setStrokeColor(colors.HexColor("#" + ORANGE))
    c.setLineWidth(2)
    c.line(20 * mm, 27 * mm, width - 20 * mm, 27 * mm)
    c.save()


def make_bundle():
    files = [
        OUT / "README.md",
        OUT / "Requerimientos-Sprint-2.docx",
        OUT / "Requerimientos-Sprint-2.pdf",
        OUT / "Supuesto-Trabajo-Integrantes-Sprint-2.docx",
        OUT / "Supuesto-Trabajo-Integrantes-Sprint-2.pdf",
        OUT / "Plan-de-Pruebas-Sprint-2.xlsx",
        OUT / "Presentacion-Sprint-2.pptx",
        DIAGRAMS / "casos-de-uso.png",
        DIAGRAMS / "modelo-de-datos.png",
        FIXTURES / "certificado-sintetico-demo.pdf",
    ]
    with ZipFile(OUT / "Entrega-RiTech-Sprint-2.zip", "w", ZIP_DEFLATED) as archive:
        for file in files:
            archive.write(file, file.relative_to(OUT))


def build():
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    FIXTURES.mkdir(parents=True, exist_ok=True)
    draw_case_diagram(DIAGRAMS / "casos-de-uso.png")
    draw_data_model(DIAGRAMS / "modelo-de-datos.png")
    make_requirements_docx()
    make_requirements_pdf()
    make_contributions_docx()
    make_contributions_pdf()
    fill_test_workbook()
    make_synthetic_certificate()
    make_bundle()
    print("Entregables generados:")
    for name in ["Requerimientos-Sprint-2.docx", "Requerimientos-Sprint-2.pdf", "Supuesto-Trabajo-Integrantes-Sprint-2.docx", "Supuesto-Trabajo-Integrantes-Sprint-2.pdf", "Plan-de-Pruebas-Sprint-2.xlsx"]:
        print(f"- {OUT / name}")
    print(f"- {DIAGRAMS / 'casos-de-uso.png'}")
    print(f"- {DIAGRAMS / 'modelo-de-datos.png'}")
    print(f"- {FIXTURES / 'certificado-sintetico-demo.pdf'}")
    print(f"- {OUT / 'Entrega-RiTech-Sprint-2.zip'}")


if __name__ == "__main__":
    build()
